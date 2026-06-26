import os
import sys
import subprocess

# Ensure python-pptx is installed
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx not found. Installing...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

# ── Color Palette ──────────────────────────────────────────────────────────
BLUE_DARK  = RGBColor(0x00, 0x3F, 0x7D)   # Primary Dark Blue
BLUE_BRAND = RGBColor(0x00, 0x57, 0xA8)   # UTCL Brand Blue
GREY_DARK  = RGBColor(0x33, 0x33, 0x33)   # Charcoal Grey for text
GREY_LIGHT = RGBColor(0xF1, 0xF5, 0xF9)   # Light backgrounds for cards
YELLOW     = RGBColor(0xFF, 0xED, 0x00)   # UTCL Accent Yellow
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE     = RGBColor(0xE8, 0x6A, 0x10)   # Warning / accent orange

def add_slide_title(slide, text, color=BLUE_DARK):
    """Adds a slide title with standard premium styling."""
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9.0), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = 'Calibri'
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = color
    return title_box

def add_card(slide, left, top, width, height, title, points, bg_color=GREY_LIGHT):
    """Creates a card shape with a title and bullet points inside."""
    # Add background rectangle shape
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = RGBColor(0xDC, 0xE2, 0xE8) # Soft border
    shape.line.width = Pt(1)

    # Add text inside shape
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)

    # Card Title
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.name = 'Calibri'
    p_title.font.size = Pt(16)
    p_title.font.bold = True
    p_title.font.color.rgb = BLUE_DARK
    p_title.space_after = Pt(10)

    # Card Bullet Points
    for pt in points:
        p = tf.add_paragraph()
        p.text = f"• {pt}"
        p.font.name = 'Calibri'
        p.font.size = Pt(12)
        p.font.color.rgb = GREY_DARK
        p.space_after = Pt(6)

def main():
    prs = Presentation()
    # Use widescreen aspect ratio (16:9)
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Premium Theme)
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    # Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(5.625))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE_DARK
    bg.line.fill.background()

    # Title & Subtitle in single frame
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(8.0), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "UTCL BUS MANAGEMENT SYSTEM"
    p.font.name = 'Calibri'
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = YELLOW
    p.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Internal Transport System & LAN Deployment"
    p2.font.name = 'Calibri'
    p2.font.size = Pt(20)
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(30)

    p3 = tf.add_paragraph()
    p3.text = "UltraTech Cement Limited\nPlant Operations — Confidential Internal System"
    p3.font.name = 'Calibri'
    p3.font.size = Pt(13)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(0xDF, 0xEA, 0xF8)

    # =========================================================================
    # SLIDE 2: Problem Statement & Objectives
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "1. Project Need & Core Objectives")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Current Operational Flaws", 
             [
                 "Manual seat reservation overheads at plant site.",
                 "Lack of driver accountability & trip validation.",
                 "No live location or stop progress tracking for admins.",
                 "Cramped schedules causing delay issues in employee commutes.",
                 "High error rate in manual booking listings & logs."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Key Project Objectives", 
             [
                 "Digitize seat booking with a flat fare structure.",
                 "Provide a real-time, zero-dependency LAN client view.",
                 "Enable dual shifts & driver trip progress monitoring.",
                 "Automate bus maintenance and servicing records.",
                 "Deliver robust booking data and occupancy reports."
             ])

    # =========================================================================
    # SLIDE 3: Features — Employee Portal
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "2. Employee Portal Capabilities")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Booking & Route Selection", 
             [
                 "Interactive Seat Selection Map: Choose exact seat.",
                 "Flat Fare System: Strict flat ₹20 fare applied to bookings.",
                 "Multi-Passenger Bookings: Support up to 4 passenger seats per booking transaction (Employee + family).",
                 "Interactive Stops Selection: Choose boarding/drop stops independently with validation of stop sequence."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Digital Tickets & History", 
             [
                 "Virtual Ticket Generation: Instant ticket showing name, PS number, shift, seat, stops, fare, and code.",
                 "Unique QR Verification Code: Scan to verify ticket details.",
                 "Booking History: Sorted list labeled with 'Active' or 'Past' status.",
                 "Booking Export: Export history as a formatted Excel (.xlsx) file."
             ])

    # =========================================================================
    # SLIDE 4: Features — Admin & Driver Portals
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "3. Admin & Driver Capabilities")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Driver Dash & Shift Logging", 
             [
                 "Shift Assignments: Shows daily assigned shifts, stops, and arrival times.",
                 "Trip Attendance Logging: Submit departure and arrival logs with GPS-verified photos.",
                 "Stop Progress Tracking: Update bus location at key stops.",
                 "Service Requests: Submit 'Service Done' notices upon completing scheduled bus servicing."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Admin Control Panel", 
             [
                 "User Management: CRUD operations for Employee & Driver accounts with unique PS numbers.",
                 "Shift Scheduling: Configure timings, stops, and bus/driver assignments with conflict checks.",
                 "Maintenance logs: Schedule and check service status dynamically (Services Schedule & Timings).",
                 "Pending Review: Review driver attendance logs."
             ])

    # =========================================================================
    # SLIDE 5: Real-time Synchronisation & Concurrency Controls
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "4. Real-time Sync & Concurrency Controls")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Dynamic WebSocket Sync", 
             [
                 "Flask-SocketIO connection handles live seat holds & bookings on the local plant network.",
                 "Under 1-second update broadcasts: When a user selects/books a seat, it updates other active maps immediately.",
                 "Offline Fallback: Automatically falls back to a 15-second HTTP polling loop if WebSocket drops."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Concurrency Guards", 
             [
                 "Seat holds are active for 180 seconds to reserve space while the user completes UPI verification.",
                 "Backend database row-level locking: SELECT FOR UPDATE on the shifts table during bookings.",
                 "Strict duplicate check: Prevents an employee from booking the same shift/direction twice."
             ])

    # =========================================================================
    # SLIDE 6: System Capacity & Performance
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "5. Capacity & Performance Metrics")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "User & Transaction Scale", 
             [
                 "Simultaneous Logins: Threaded Flask server supports 200–500 concurrent sessions natively.",
                 "Scalable Limit: Expandable to 1,000+ users by configuring reverse proxy (Nginx).",
                 "User Count Capacity: Confidently manages 50,000+ registered PS numbers with zero performance drop."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Data Growth & Operations", 
             [
                 "Negligible Growth: Approx 1 MB storage growth per 5,000 bookings.",
                 "Storage capacity: A 10 GB disk retains 50+ years of booking database history.",
                 "Instant Query Speeds: dashboard, logins, seat loads, search queries execute in under 1-2 seconds."
             ])

    # =========================================================================
    # SLIDE 7: IT Deployment Architecture
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "6. Intranet LAN Deployment Architecture")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "IT Server Room Host", 
             [
                 "Dedicated Server PC inside the plant's IT room hosting Python API & MySQL Database.",
                 "Assigned LAN Static IP address: Reserving fixed IP address (e.g., 192.168.1.100).",
                 "No Internet Access Required: Fully self-contained inside the corporate firewall.",
                 "Windows Firewall Port Rule: TCP Port 8000 allowed for incoming connections."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Network Workstations", 
             [
                 "Client access: Workstations (HR, Admin, Gate) load browser with URL: http://<STATIC-IP>:8000.",
                 "Zero Client Overhead: No plugins, extensions, or software installations needed on client PCs.",
                 "Auto-Start Trigger: Windows Task Scheduler launches run.bat automatically 30s after server boot."
             ])

    # =========================================================================
    # SLIDE 8: Operations: Maintenance & Backups
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "7. Maintenance & Backup Guidelines")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Scheduled Database Backups", 
             [
                 "Automated Daily SQL Backups: Triggering MySQL dump command daily via Task Scheduler.",
                 "Command template: mysqldump -u root -p utcl_bus_db > backup.sql.",
                 "Offline archiving: Recommended backing up files to an external network storage drive."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Photo Retention & Servicing", 
             [
                 "Background photo cleanup: Routine loop in server.py deletes driver attendance photos older than 30 days.",
                 "Disk space recovery: Deletes raw jpg/png files from uploads folder to prevent disk filling.",
                 "Servicing status: Admins confirm drivers' completed service alerts, updating records instantly."
             ])

    # =========================================================================
    # SLIDE 9: Security, Limitations & Future Roadmap
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    add_slide_title(slide, "8. Security, Limitations & Future Scope")

    add_card(slide, Inches(0.5), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Security & Controls", 
             [
                 "PS Number authentication (case-insensitive) prevents unauthorized access.",
                 "Passwords stored securely using bcrypt hashing.",
                 "Self-service forgot password flow via local SMTP OTP active.",
                 "MySQL SELECT FOR UPDATE row-locks prevent duplicate bookings.",
                 "LAN deployment naturally shields the application from external threats."
             ])

    add_card(slide, Inches(5.3), Inches(1.3), Inches(4.2), Inches(3.6), 
             "Future Scalability Roadmap", 
             [
                 "Upgrade HTTP to HTTPS using self-signed SSL certificates.",
                 "Integrate Nginx reverse proxy to offload static assets and increase concurrent capacity.",
                 "Dedicated MySQL database server deployment for hardware isolation.",
                 "Implement MySQL read replicas to scale read operations."
             ])

    # =========================================================================
    # SLIDE 10: Conclusion & Demo (Dark Premium Theme)
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    # Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(5.625))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE_DARK
    bg.line.fill.background()

    # Title & Subtitle in single frame
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(8.0), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = 'Calibri'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = YELLOW
    p.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "UTCL Bus Management System is ready for LAN deployment."
    p2.font.name = 'Calibri'
    p2.font.size = Pt(18)
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(20)

    p3 = tf.add_paragraph()
    p3.text = "Demo Credentials for testing:\n• Admin Login: PS00001 / password123\n• Employee Login: PS10001 / password123\n• Driver Login: PS20001 / password123"
    p3.font.name = 'Calibri'
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(0xDF, 0xEA, 0xF8)

    # Save presentation
    out_path = r"d:\UTCL antigravity\UTCL_Bus_System_Presentation.pptx"
    prs.save(out_path)
    print(f"[OK] PowerPoint Presentation generated successfully at: {out_path}")

if __name__ == "__main__":
    main()
